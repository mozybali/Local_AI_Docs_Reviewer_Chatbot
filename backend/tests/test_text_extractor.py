"""Metin çıkarma testleri.

- Sanitizasyon: bazı PDF'ler glifleri NUL (0x00) baytına eşler; bu metin
  PostgreSQL'e yazılırken `DataError` verir. `extract_pages` çıkışta NUL ve
  diğer kontrol karakterlerini temizlemeli.
- OCR fallback: native metni yetersiz sayfalar OCR'a gider; OCR yoksa/başarısız
  olursa sayfa `failed`, gerçekten boş sayfa `empty` işaretlenir. OCR motoru
  sahte (fake) ile taklit edilir; gerçek motor/model gerekmez.
- DOCX: paragrafların yanında tablolar da (satır bütünlüğüyle) çıkarılmalı.
- Görüntü dosyaları: OCR ile okunur; OCR yoksa açıklayıcı hata verilir.
"""

import pytest

from app.services import ocr_service, text_extractor
from app.services.ocr_service import OCRResult
from app.services.text_extractor import (
    TextExtractionError,
    _RawPage,
    _sanitize_text,
    extract_pages,
)


def test_sanitize_removes_nul_and_control_chars():
    # NUL ve diğer C0 kontrol karakterleri kaldırılır; harf/boşluk korunur.
    dirty = "Ün\x00vers\x00tes\x00\x07 \tsatır\nsonu\r"
    clean = _sanitize_text(dirty)
    assert "\x00" not in clean
    assert "\x07" not in clean
    assert clean.startswith("Ünverstes")  # NUL'lar düştü
    # Yaygın boşluklar (tab, newline, carriage return) korunur.
    assert "\t" in clean and "\n" in clean and "\r" in clean


def test_sanitize_handles_empty():
    assert _sanitize_text("") == ""


def test_extract_txt_strips_nul_bytes(tmp_path):
    # PDF'teki NUL-glif durumunu TXT üzerinden taklit et: çıkış NUL içermemeli.
    f = tmp_path / "belge.txt"
    f.write_bytes("Ün\x00vers\x00te metni".encode("utf-8"))

    pages = extract_pages(str(f), "txt")
    assert len(pages) == 1
    assert "\x00" not in pages[0].text
    assert pages[0].text == "Ünverste metni"


# --- PDF OCR fallback ------------------------------------------------------


NATIVE_TEXT = (
    "Bu sayfanın metin katmanı gayet yeterli uzunlukta bir içerik taşıyor."
)


def _patch_native_pdf(monkeypatch, raw_pages):
    """Native PDF çıkarıcılarını sahte sonuçla değiştirir."""
    monkeypatch.setattr(
        text_extractor,
        "_extract_pdf_with_pdfplumber",
        lambda path: raw_pages,
    )
    monkeypatch.setattr(
        text_extractor, "_extract_pdf_with_pypdf", lambda path: None
    )


def test_scanned_page_falls_back_to_ocr(monkeypatch):
    """Metinsiz + görüntülü sayfa OCR ile okunmalı, kaynağı `ocr` olmalı."""
    _patch_native_pdf(monkeypatch, [
        _RawPage(page_number=1, text=NATIVE_TEXT, has_images=False),
        _RawPage(page_number=2, text="", has_images=True),
    ])
    monkeypatch.setattr(text_extractor.ocr_service, "is_available", lambda: True)
    monkeypatch.setattr(
        text_extractor.ocr_service,
        "ocr_pdf_page",
        lambda path, idx: OCRResult(
            text="Taranmış sayfadan OCR ile okunan metin.",
            confidence=0.92,
            engine="fake",
        ),
    )

    pages = extract_pages("sahte.pdf", "pdf")
    assert pages[0].source == "native"
    assert pages[0].text == NATIVE_TEXT
    assert pages[1].source == "ocr"
    assert pages[1].ocr_confidence == 0.92
    assert "OCR ile okunan" in pages[1].text


def test_scanned_page_without_ocr_marked_failed(monkeypatch):
    """OCR kullanılamıyorsa görüntülü-metinsiz sayfa `failed` işaretlenmeli."""
    _patch_native_pdf(monkeypatch, [
        _RawPage(page_number=1, text="", has_images=True),
    ])
    monkeypatch.setattr(
        text_extractor.ocr_service, "is_available", lambda: False
    )

    pages = extract_pages("sahte.pdf", "pdf")
    assert pages[0].source == "failed"
    assert pages[0].text == ""


def test_low_confidence_ocr_marked_failed(monkeypatch):
    """Güven eşiğinin altındaki OCR sonucu çöp sayılır; sayfa `failed` olur."""
    _patch_native_pdf(monkeypatch, [
        _RawPage(page_number=1, text="", has_images=True),
    ])
    monkeypatch.setattr(text_extractor.ocr_service, "is_available", lambda: True)
    monkeypatch.setattr(
        text_extractor.ocr_service,
        "ocr_pdf_page",
        lambda path, idx: OCRResult(text="çöp", confidence=0.05, engine="fake"),
    )

    pages = extract_pages("sahte.pdf", "pdf")
    assert pages[0].source == "failed"


def test_truly_blank_page_marked_empty_without_ocr_call(monkeypatch):
    """Görüntüsüz + metinsiz sayfa OCR'a hiç gitmeden `empty` olmalı."""
    _patch_native_pdf(monkeypatch, [
        _RawPage(page_number=1, text="", has_images=False),
    ])
    monkeypatch.setattr(text_extractor.ocr_service, "is_available", lambda: True)

    def boom(path, idx):  # OCR çağrılırsa test patlasın
        raise AssertionError("Boş sayfa için OCR çağrılmamalıydı")

    monkeypatch.setattr(text_extractor.ocr_service, "ocr_pdf_page", boom)

    pages = extract_pages("sahte.pdf", "pdf")
    assert pages[0].source == "empty"


# --- DOCX tabloları ---------------------------------------------------------


def test_docx_tables_are_extracted(tmp_path):
    """DOCX içindeki tablo hücreleri satır bütünlüğüyle metne dahil edilmeli."""
    import docx

    d = docx.Document()
    d.add_paragraph("Personel listesi aşağıdadır.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Ad"
    table.cell(0, 1).text = "Departman"
    table.cell(1, 0).text = "Ayşe"
    table.cell(1, 1).text = "Muhasebe"
    path = tmp_path / "tablolu.docx"
    d.save(str(path))

    pages = extract_pages(str(path), "docx")
    assert len(pages) == 1
    text = pages[0].text
    assert "Personel listesi" in text
    assert "Ad | Departman" in text
    assert "Ayşe | Muhasebe" in text


# --- Görüntü dosyaları ------------------------------------------------------


def test_image_file_uses_ocr(monkeypatch, tmp_path):
    monkeypatch.setattr(text_extractor.ocr_service, "is_available", lambda: True)
    monkeypatch.setattr(
        text_extractor.ocr_service,
        "ocr_image_file",
        lambda path: OCRResult(
            text="Görüntüdeki fatura metni.", confidence=0.88, engine="fake"
        ),
    )

    pages = extract_pages(str(tmp_path / "fatura.png"), "png")
    assert len(pages) == 1
    assert pages[0].source == "ocr"
    assert pages[0].ocr_confidence == 0.88
    assert "fatura metni" in pages[0].text


def test_image_file_without_ocr_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(
        text_extractor.ocr_service, "is_available", lambda: False
    )
    with pytest.raises(TextExtractionError, match="OCR"):
        extract_pages(str(tmp_path / "fatura.png"), "png")
