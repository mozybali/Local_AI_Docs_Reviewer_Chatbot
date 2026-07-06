"""Metin çıkarma servisi (PDF, TXT, DOCX ve görüntü dosyaları).

Sayfa numarası bilgisini koruyacak şekilde, doküman içeriğini sayfa sayfa
çıkarır ve her sayfa için içeriğin NEREDEN geldiğini işaretler:

- ``native``: dosyanın kendi metin katmanından çıkarıldı,
- ``ocr``   : metin katmanı yetersizdi; sayfa görüntüye çevrilip lokal OCR ile
  okundu (`ocr_service`, tamamen lokal — API çağrısı yok),
- ``empty`` : sayfa gerçekten boş (metin de görüntü de yok),
- ``failed``: sayfada görüntü var ama metin çıkarılamadı (OCR kapalı/kurulu
  değil ya da OCR güveni eşiğin altında). Bu sayfalar kullanıcıya raporlanır.

PDF akışı: önce `pdfplumber` (tablo + çok kolon farkındalıklı), o tümüyle
başarısız olursa `pypdf`. pdfplumber "başarılı" görünüp hiç metin döndürmezse
(bozuk metin katmanı) pypdf ayrıca kalite tetiklemeli olarak denenir. Sayfa
bazında yetersiz metin kalırsa OCR fallback devreye girer.

DOCX: paragraflar VE tablolar belge sırasıyla okunur (python-docx
`iter_inner_content`); tablo satırları " | " ile birleştirilir. Sayfa kavramı
olmadığı için tüm içerik tek sayfadır.

Görüntü dosyaları (png/jpg/jpeg): doğrudan OCR'dan geçirilir.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from app.config import settings
from app.services import ocr_service

logger = logging.getLogger(__name__)


class TextExtractionError(Exception):
    """Metin çıkarma sırasında oluşan hata."""


# Sayfa içeriğinin kaynağı (ExtractedPage.source değerleri).
SOURCE_NATIVE = "native"
SOURCE_OCR = "ocr"
SOURCE_EMPTY = "empty"
SOURCE_FAILED = "failed"

_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg"}

# PostgreSQL `text`/`varchar` alanları NUL (0x00) baytını saklayamaz. Bazı
# PDF'ler (ör. e-imzalı Resmî Gazete belgeleri) bazı glifleri 0x00'a eşler; bu
# da chunk_text PG'ye yazılırken `DataError`'a yol açar. Bu yüzden NUL'u ve diğer
# C0 kontrol karakterlerini (yaygın boşluklar \t \n \r hariç) çıkarımda temizle.
# DEL (0x7F) de güvenlik için kaldırılır.
_CONTROL_CHARS_TABLE = {
    codepoint: None
    for codepoint in (*range(0x20), 0x7F)
    if codepoint not in (0x09, 0x0A, 0x0D)
}


def _sanitize_text(text: str) -> str:
    """Saklanamayan kontrol karakterlerini (NUL dahil) metinden temizler."""
    if not text:
        return text
    return text.translate(_CONTROL_CHARS_TABLE)


@dataclass(frozen=True)
class ExtractedPage:
    """Tek bir sayfanın çıkarılmış metni ve çıkarım meta bilgisi."""

    page_number: int
    text: str
    source: str = SOURCE_NATIVE
    ocr_confidence: float | None = None


@dataclass(frozen=True)
class _RawPage:
    """Native çıkarımdan gelen ara sayfa temsili.

    `has_images` üç durumludur: True/False (pdfplumber bilir) veya None
    (pypdf yolunda bilinmez; OCR adaylığında "olabilir" kabul edilir).
    """

    page_number: int
    text: str
    has_images: bool | None


def extract_pages(file_path: str, file_type: str) -> list[ExtractedPage]:
    """Dosya tipine göre metni sayfa sayfa çıkarır.

    Desteklenmeyen tip için `TextExtractionError` fırlatır.
    """
    ext = file_type.strip().lower().lstrip(".")
    if ext == "pdf":
        pages = _extract_pdf(file_path)
    elif ext == "txt":
        pages = _extract_txt(file_path)
    elif ext == "docx":
        pages = _extract_docx(file_path)
    elif ext in _IMAGE_EXTENSIONS:
        pages = _extract_image(file_path)
    else:
        raise TextExtractionError(f"Desteklenmeyen dosya tipi: {ext}")

    # Tüm çıkarıcılar için tek noktada sanitizasyon: PG'nin reddettiği NUL ve
    # diğer kontrol karakterlerini temizle (yoksa commit DataError verir).
    return [
        ExtractedPage(
            page_number=p.page_number,
            text=_sanitize_text(p.text),
            source=p.source,
            ocr_confidence=p.ocr_confidence,
        )
        for p in pages
    ]


def extract_text(file_path: str, file_type: str) -> str:
    """Tüm sayfaları birleştirip tek bir metin döner."""
    pages = extract_pages(file_path, file_type)
    return "\n\n".join(page.text for page in pages if page.text).strip()


# --- TXT -------------------------------------------------------------------


def _extract_txt(file_path: str) -> list[ExtractedPage]:
    """TXT dosyasını okur. UTF-8 başarısız olursa Latin-1'e düşer."""
    raw = Path(file_path).read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1", errors="replace")
    return [ExtractedPage(page_number=1, text=text.strip())]


# --- DOCX ------------------------------------------------------------------


def _format_table_rows(rows: list[list[str | None]]) -> str:
    """Tablo satırlarını, hücreleri " | " ile ayırarak metinleştirir.

    Retrieval'da "satır bütünlüğü" korunur: bir hücre değeri, aynı satırdaki
    diğer hücrelerle (ör. etiketiyle) yan yana kalır.
    """
    lines: list[str] = []
    for row in rows:
        cells = [(cell or "").strip().replace("\n", " ") for cell in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_docx(file_path: str) -> list[ExtractedPage]:
    """DOCX dosyasından paragrafları VE tabloları belge sırasıyla çıkarır.

    python-docx sayfa sınırlarını vermediği için tüm içerik tek bir sayfa
    olarak birleştirilir. Bozuk/desteklenmeyen dosyalar `TextExtractionError`
    ile işaretlenir.
    """
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:  # pragma: no cover - bağımlılık eksikse
        raise TextExtractionError(
            "DOCX desteği için python-docx kurulu değil."
        )

    try:
        document = docx.Document(file_path)
    except Exception:
        raise TextExtractionError(
            "Bu DOCX dosyası okunamadı. Bozuk olabilir."
        )

    # Belge sırasını koruyarak paragraf + tablo dolaş (python-docx >= 1.1).
    # Eski sürümlerde sıra korunamaz; paragraflar sonra tablolar okunur.
    try:
        blocks = list(document.iter_inner_content())
    except AttributeError:  # pragma: no cover - eski python-docx
        blocks = [*document.paragraphs, *document.tables]

    parts: list[str] = []
    for block in blocks:
        if isinstance(block, Paragraph):
            if block.text and block.text.strip():
                parts.append(block.text.strip())
        elif isinstance(block, Table):
            rows = [
                [cell.text for cell in row.cells] for row in block.rows
            ]
            table_text = _format_table_rows(rows)
            if table_text:
                parts.append(table_text)

    text = "\n".join(parts).strip()
    return [ExtractedPage(page_number=1, text=text)]


# --- Görüntü dosyaları ------------------------------------------------------


def _extract_image(file_path: str) -> list[ExtractedPage]:
    """Görüntü dosyasını (png/jpg/jpeg) OCR ile metne çevirir."""
    if not settings.OCR_ENABLED or not ocr_service.is_available():
        raise TextExtractionError(
            "Görüntü dosyalarını işlemek için OCR gerekir ancak hiçbir OCR "
            "motoru kullanılamıyor. Tesseract kurun veya `pip install easyocr` "
            "çalıştırın."
        )

    try:
        result = ocr_service.ocr_image_file(file_path)
    except ocr_service.OCRError as exc:
        raise TextExtractionError(f"Görüntü OCR ile okunamadı: {exc}")

    if result is None or not result.text.strip():
        raise TextExtractionError(
            "Görüntüden metin çıkarılamadı. Görüntü metin içermiyor veya "
            "kalite çok düşük olabilir."
        )
    if (
        result.confidence is not None
        and result.confidence < settings.OCR_MIN_CONFIDENCE
    ):
        raise TextExtractionError(
            "Görüntüdeki metin güvenilir şekilde okunamadı (düşük OCR "
            f"güveni: {result.confidence:.2f}). Daha yüksek çözünürlüklü bir "
            "görüntü deneyin."
        )
    return [
        ExtractedPage(
            page_number=1,
            text=result.text.strip(),
            source=SOURCE_OCR,
            ocr_confidence=result.confidence,
        )
    ]


# --- PDF ---------------------------------------------------------------------


def _extract_pdf(file_path: str) -> list[ExtractedPage]:
    """PDF metnini sayfa sayfa çıkarır (native + kalite kontrol + OCR fallback)."""
    raw_pages = _extract_pdf_with_pdfplumber(file_path)
    if raw_pages is None:
        raw_pages = _extract_pdf_with_pypdf(file_path)

    if raw_pages is None:
        raise TextExtractionError(
            "Bu PDF okunamadı. Şifreli veya bozuk olabilir."
        )

    # Kalite tetiklemeli ikinci deneme: pdfplumber exception atmadı ama hiçbir
    # sayfadan kayda değer metin gelmediyse (bozuk metin katmanı) pypdf'i de
    # dene; sayfa bazında daha uzun olan metni kullan.
    threshold = settings.OCR_TRIGGER_MIN_CHARS
    if raw_pages and all(len(p.text) < threshold for p in raw_pages):
        alternative = _extract_pdf_with_pypdf(file_path)
        if alternative and len(alternative) == len(raw_pages):
            raw_pages = [
                alt if len(alt.text) > len(orig.text) else orig
                for orig, alt in zip(raw_pages, alternative)
            ]

    return [_finalize_pdf_page(file_path, raw) for raw in raw_pages]


def _finalize_pdf_page(file_path: str, raw: _RawPage) -> ExtractedPage:
    """Native sonucu kalite kontrolünden geçirir; gerekirse OCR fallback uygular.

    Karar akışı:
    1. Native metin eşik üstündeyse sayfa `native` kabul edilir.
    2. Eşik altındaysa ve sayfada görüntü olabilir ise OCR denenir.
    3. OCR sonucu güven eşiğini geçerse `ocr`; geçemezse eldeki native kırıntı
       korunur; hiçbir şey yoksa sayfa `failed` (görüntü varken) veya `empty`
       (gerçekten boş sayfa) olur.
    """
    text = raw.text.strip()
    if len(text) >= settings.OCR_TRIGGER_MIN_CHARS:
        return ExtractedPage(
            page_number=raw.page_number, text=text, source=SOURCE_NATIVE
        )

    # Görüntüsü olmadığı KESİN bilinen sayfada OCR anlamsızdır (boş sayfa).
    ocr_possible = (
        settings.OCR_ENABLED
        and raw.has_images is not False
        and ocr_service.is_available()
    )

    if ocr_possible:
        result = None
        try:
            result = ocr_service.ocr_pdf_page(file_path, raw.page_number - 1)
        except ocr_service.OCRError as exc:
            logger.warning(
                "OCR hatası (sayfa %s): %s", raw.page_number, exc
            )

        if result is not None and result.text.strip():
            confident = (
                result.confidence is None
                or result.confidence >= settings.OCR_MIN_CONFIDENCE
            )
            if confident:
                return ExtractedPage(
                    page_number=raw.page_number,
                    text=result.text.strip(),
                    source=SOURCE_OCR,
                    ocr_confidence=result.confidence,
                )
            logger.info(
                "OCR güveni eşiğin altında (sayfa %s): %.2f < %.2f",
                raw.page_number,
                result.confidence,
                settings.OCR_MIN_CONFIDENCE,
            )

    # OCR yapılamadı/başarısız: elde native kırıntı varsa onu koru.
    if text:
        return ExtractedPage(
            page_number=raw.page_number, text=text, source=SOURCE_NATIVE
        )
    if raw.has_images is not False:
        # Sayfada görüntü var(dı) ama metin çıkarılamadı: kullanıcıya
        # raporlanacak gerçek bir kayıp.
        return ExtractedPage(
            page_number=raw.page_number, text="", source=SOURCE_FAILED
        )
    return ExtractedPage(page_number=raw.page_number, text="", source=SOURCE_EMPTY)


def _extract_pdf_with_pdfplumber(file_path: str) -> list[_RawPage] | None:
    """pdfplumber ile dener; başarısız olursa None döner.

    Sayfa metni tablo ve çok kolon farkındalıklı çıkarılır
    (`_page_text_with_layout`); ayrıca sayfada görüntü olup olmadığı OCR
    adaylığı için işaretlenir.
    """
    try:
        import pdfplumber
    except ImportError:
        return None

    try:
        pages: list[_RawPage] = []
        with pdfplumber.open(file_path) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = _page_text_with_layout(page)
                has_images = bool(page.images)
                pages.append(
                    _RawPage(
                        page_number=index,
                        text=text.strip(),
                        has_images=has_images,
                    )
                )
                # pdfplumber sayfa nesneleri layout cache'i biriktirir; büyük
                # PDF'lerde bellek sayfa sayısıyla büyür. Metni aldıktan sonra
                # cache bırakılır (sayfa başına sabit bellek).
                page.flush_cache()
        return pages
    except Exception:
        return None


def _page_text_with_layout(page) -> str:
    """Bir pdfplumber sayfasından layout farkındalıklı metin çıkarır.

    - Tablolar tespit edilirse: tablo alanları gövde metninden çıkarılır ve
      tablolar satır bütünlüğü korunarak (" | " ayraçlı) sona eklenir.
      Böylece tablo hücreleri okuma sırasına saçılmaz.
    - Tablo yoksa iki kolonlu düzen aranır; bulunursa kolonlar ayrı ayrı
      (önce sol, sonra sağ) çıkarılır ve okuma sırası korunur.
    - Her adım muhafazakârdır: emin olunamayan durumda düz `extract_text()`
      davranışına düşülür (mevcut davranış asla kötüleşmez).
    """
    try:
        tables = page.find_tables()
    except Exception:
        tables = []

    if tables:
        try:
            return _text_with_tables(page, tables)
        except Exception:
            logger.debug(
                "Tablo farkındalıklı çıkarım başarısız; düz metne dönülüyor.",
                exc_info=True,
            )

    try:
        column_text = _extract_two_columns(page)
        if column_text:
            return column_text
    except Exception:
        logger.debug(
            "Kolon tespiti başarısız; düz metne dönülüyor.", exc_info=True
        )

    return page.extract_text() or ""


def _text_with_tables(page, tables) -> str:
    """Tablo alanlarını gövdeden ayırıp tablo metnini yapılandırılmış ekler."""
    bboxes = [t.bbox for t in tables]

    def outside_tables(obj) -> bool:
        v_mid = (obj["top"] + obj["bottom"]) / 2
        h_mid = (obj["x0"] + obj["x1"]) / 2
        return not any(
            x0 <= h_mid <= x1 and top <= v_mid <= bottom
            for (x0, top, x1, bottom) in bboxes
        )

    body = page.filter(outside_tables).extract_text() or ""
    table_parts = []
    for table in tables:
        rows = table.extract()
        formatted = _format_table_rows(rows or [])
        if formatted:
            table_parts.append(formatted)

    sections = [s for s in (body.strip(), *table_parts) if s]
    return "\n\n".join(sections)


def _extract_two_columns(page) -> str | None:
    """İki kolonlu sayfa düzenini tespit edip kolonları sırayla çıkarır.

    Sayfanın orta bandında (genişliğin %38-62'si) hiçbir kelimenin kesmediği
    dikey bir boşluk (gutter) aranır. Bulunursa ve her iki taraf da anlamlı
    miktarda kelime içeriyorsa, sol ve sağ kolonlar ayrı ayrı çıkarılır.
    Emin olunamayan her durumda None döner (düz çıkarıma düşülür).
    """
    words = page.extract_words()
    if len(words) < 30:
        return None

    width = float(page.width)
    gutter = None
    for percent in range(38, 63):
        candidate = width * percent / 100.0
        crossing = sum(1 for w in words if w["x0"] < candidate < w["x1"])
        if crossing:
            continue
        left_count = sum(1 for w in words if w["x1"] <= candidate)
        right_count = len(words) - left_count
        if min(left_count, right_count) >= max(5, int(0.25 * len(words))):
            gutter = candidate
            break

    if gutter is None:
        return None

    left = page.crop((0, 0, gutter, page.height)).extract_text() or ""
    right = page.crop((gutter, 0, page.width, page.height)).extract_text() or ""
    combined = "\n".join(s for s in (left.strip(), right.strip()) if s)
    return combined or None


def _extract_pdf_with_pypdf(file_path: str) -> list[_RawPage] | None:
    """pypdf ile dener; başarısız olursa None döner.

    pypdf sayfadaki görüntü varlığını ucuz şekilde bildirmediği için
    `has_images=None` (bilinmiyor) işaretlenir; OCR adaylığı dışlanmaz.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return None

    try:
        reader = PdfReader(file_path)
        if reader.is_encrypted:
            # Boş şifre ile açmayı dene; olmuyorsa başarısız say.
            try:
                reader.decrypt("")
            except Exception:
                return None
        pages: list[_RawPage] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(
                _RawPage(page_number=index, text=text.strip(), has_images=None)
            )
        return pages
    except Exception:
        return None
